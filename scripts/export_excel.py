import json, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "data")
SITE = os.path.join(ROOT, "docs")
LINK_FONT = Font(color="0563C1", underline="single")
HEAD_FILL = PatternFill("solid", fgColor="1F2937")
HEAD_FONT = Font(color="FFFFFF", bold=True)

def build(repos, path, cols):
    wb = Workbook()
    ws = wb.active
    ws.title = "ranking"
    ws.append([c[0] for c in cols])
    for cell in ws[1]:
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
    for x in repos:
        ws.append([x.get(key, "") for _, key in cols])
    link_idx = [i for i, c in enumerate(cols, start=1) if c[1] == "url"][0]
    for r in range(2, ws.max_row + 1):
        cell = ws.cell(row=r, column=link_idx)
        cell.hyperlink = cell.value
        cell.font = LINK_FONT
    widths = {"rank": 8, "full_name": 34, "descripcion_es": 80, "description": 70,
              "stars": 11, "forks": 13, "language": 13, "categoria": 19,
              "license": 12, "tipo": 17, "pushed_at": 12, "created_at": 12,
              "archived": 11, "url": 44}
    for i, c in enumerate(cols, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = widths.get(c[1], 16)
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"
    try:
        wb.save(path)
    except PermissionError:
        alt = path.replace(".xlsx", "-nuevo.xlsx")
        wb.save(alt)
        path = alt
    print("ok", os.path.basename(path), ws.max_row - 1, "filas")

top = json.load(open(os.path.join(SRC, "repos.json"), encoding="utf-8"))["repos"]
trend = json.load(open(os.path.join(SRC, "trending.json"), encoding="utf-8"))["repos"]

BASE_COLS = [("puesto", "rank"), ("nombre", "full_name"), ("descripcion", "descripcion_es"),
             ("descripcion_original", "description"), ("estrellas", "stars"),
             ("bifurcaciones", "forks"), ("lenguaje", "language"), ("categoria", "categoria"),
             ("licencia", "license"), ("tipo", "tipo")]
build(top, os.path.join(SITE, "top500.xlsx"),
      BASE_COLS + [("actividad", "pushed_at"), ("archivado", "archived"), ("enlace", "url")])
build(trend, os.path.join(SITE, "trending.xlsx"),
      BASE_COLS + [("creado", "created_at"), ("actividad", "pushed_at"), ("enlace", "url")])
